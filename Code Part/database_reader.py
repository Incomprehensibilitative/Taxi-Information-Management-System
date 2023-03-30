"""Reader funcion:
Read all data in a sheet and return in an object list"""

import openpyxl

from Domains import Customer, Vehicle, Driver, Invoice

# read database
data = openpyxl.load_workbook("Taxi-information.xlsx", read_only=True, data_only=True)


def take_customer_info():
    sheet = data['Customer']
    customer_list = []
    for row in sheet.iter_rows(min_row=2):
        new_customer = Customer.Customer()
        new_customer.set_customer(row[0].value, row[1].value, row[2].value)
        customer_list.append(new_customer)
    return customer_list


def take_vehicle_info():
    sheet = data['Vehicle']
    vehicle_list = []
    for row in sheet.iter_rows(min_row=2):
        new_vehicle = Vehicle.Vehicle()
        new_vehicle.set_vehicle(row[0].value, row[1].value, row[2].value, row[3].value)
        vehicle_list.append(new_vehicle)
    return vehicle_list


def take_driver_info():
    sheet = data['Driver']
    driver_list = []
    for row in sheet.iter_rows(min_row=2):
        new_driver = Driver.Driver()
        new_driver.set_driver(row[0].value, row[1].value, row[2].value, row[3].value, row[4].value, row[5].value,
                              row[6].value)
        driver_list.append(new_driver)
    return driver_list


def take_invoice_info():
    sheet = data['Invoice']
    invoice_list = []
    for row in sheet.iter_rows(min_row=2):
        new_invoice = Invoice.Invoice()
        new_invoice.set_invoice(row[0].value, row[1].value, row[2].value, row[3].value, row[4].value, row[5].value,
                                row[6].value, row[7].value)
        invoice_list.append(new_invoice)
    return invoice_list